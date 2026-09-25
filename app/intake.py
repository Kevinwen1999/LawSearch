"""Scenario file intake (Phase 6): PDF/DOCX/text extraction with an OCR fallback for
image-only pages, plus content-addressed storage of the original upload in MinIO.

Extracted text is never logged — only the content hash and object key are.
"""

import hashlib
import io
import logging
from dataclasses import dataclass
from functools import lru_cache

import pypdfium2 as pdfium
from docx import Document
from minio import Minio
from minio.error import S3Error

from app.config import settings

logger = logging.getLogger(__name__)

# Below this many characters, a PDF page is treated as image-only (a scanned filing, not a
# text-layer PDF) and sent to OCR instead.
MIN_TEXT_LAYER_CHARS = 20
OCR_RENDER_SCALE = 2.0  # ~144 DPI; enough for OCR without huge intermediate bitmaps


class ExtractionError(ValueError):
    pass


# RapidOCR's bundled recognizer (ch_PP-OCRv4) is trained mostly on Chinese and drops the spaces
# between English words ("Myclientfledpersecution..."), and its text-direction classifier flips
# some upright lines 180° into gibberish. On scanned renderings of the eval scenarios
# (scripts/eval_intake.py) word agreement with the typed text was 0.25; PaddleOCR's English
# recognizer with the classifier off reads 0.997. Scans are upright, so the classifier isn't needed.
OCR_REC_MODEL = ("SWHL/RapidOCR", "PP-OCRv3/en_PP-OCRv3_rec_infer.onnx")


@lru_cache(maxsize=1)
def _ocr_engine():
    from huggingface_hub import hf_hub_download
    from rapidocr_onnxruntime import RapidOCR

    try:
        rec_model = hf_hub_download(*OCR_REC_MODEL)
    except Exception as exc:  # offline and not cached: the bundled model still reads, just badly
        logger.warning("English OCR model unavailable (%s); falling back to RapidOCR's default", type(exc).__name__)
        rec_model = None
    return RapidOCR(rec_model_path=rec_model, use_cls=False) if rec_model else RapidOCR(use_cls=False)


def _ocr_page(page: pdfium.PdfPage) -> str:
    bitmap = page.render(scale=OCR_RENDER_SCALE)
    image = bitmap.to_pil().convert("RGB")
    import numpy as np

    result, _ = _ocr_engine()(np.array(image))
    if not result:
        return ""
    # RapidOCR returns [box, text, score] per detected line; box[0] is the top-left corner,
    # so sorting by its y then x approximates reading order for a simple single-column page.
    result.sort(key=lambda r: (r[0][0][1], r[0][0][0]))
    return "\n".join(line[1] for line in result)


def _extract_pdf(content: bytes) -> str:
    pages_text = []
    pdf = pdfium.PdfDocument(content)
    try:
        n_pages = min(len(pdf), settings.ocr_max_pages)
        for i in range(n_pages):
            page = pdf[i]
            try:
                text = page.get_textpage().get_text_range().strip()
                if len(text) < MIN_TEXT_LAYER_CHARS:
                    text = _ocr_page(page).strip()
            finally:
                page.close()
            pages_text.append(text)
    finally:
        pdf.close()
    return "\n\n".join(t for t in pages_text if t)


def _extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_text(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_text,
    ".md": _extract_text,
}


def extract_text(filename: str, content: bytes) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise ExtractionError(f"unsupported file type {ext or '(none)'!r}; use PDF, DOCX or text")
    text = extractor(content).strip()
    if not text:
        raise ExtractionError("no text could be extracted from this file")
    return text


def content_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


@lru_cache(maxsize=1)
def _client() -> Minio:
    return Minio(
        settings.minio_endpoint,
        access_key=settings.minio_access_key,
        secret_key=settings.minio_secret_key,
        secure=settings.minio_secure,
    )


def _ensure_bucket() -> None:
    client = _client()
    if not client.bucket_exists(settings.minio_bucket):
        client.make_bucket(settings.minio_bucket)


@dataclass
class StoredUpload:
    key: str
    hash: str


def store_upload(filename: str, content: bytes) -> StoredUpload:
    """Store the original file keyed by its content hash, so re-uploading the same file is a
    no-op and unrelated uploads never collide."""
    _ensure_bucket()
    digest = content_hash(content)
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    key = f"uploads/{digest}{ext}"
    _client().put_object(settings.minio_bucket, key, io.BytesIO(content), length=len(content))
    return StoredUpload(key=key, hash=digest)


def delete_upload(key: str) -> None:
    try:
        _client().remove_object(settings.minio_bucket, key)
    except S3Error as exc:
        if exc.code != "NoSuchKey":
            raise
