import io

import pytest
from docx import Document

from app import intake
from app.intake import ExtractionError, content_hash, extract_text


def test_txt_extraction_decodes_utf8():
    assert extract_text("scenario.txt", "café leak, six months".encode()) == "café leak, six months"


def test_md_extraction_uses_same_path_as_txt():
    assert extract_text("scenario.md", b"# heading\nfacts") == "# heading\nfacts"


def test_docx_extraction_joins_non_empty_paragraphs():
    doc = Document()
    doc.add_paragraph("The tenant's roof leaked for six months.")
    doc.add_paragraph("")  # blank paragraphs are dropped
    doc.add_paragraph("The landlord did not respond to written requests.")
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text("scenario.docx", buf.getvalue())

    assert text == (
        "The tenant's roof leaked for six months.\n"
        "The landlord did not respond to written requests."
    )


def test_unsupported_extension_raises():
    with pytest.raises(ExtractionError, match="unsupported file type"):
        extract_text("scenario.exe", b"whatever")


def test_no_extension_raises():
    with pytest.raises(ExtractionError, match="unsupported file type"):
        extract_text("scenario", b"whatever")


def test_empty_content_raises():
    with pytest.raises(ExtractionError, match="no text could be extracted"):
        extract_text("scenario.txt", b"   \n\n  ")


def test_pdf_dispatches_to_pdf_extractor(monkeypatch):
    monkeypatch.setitem(intake._EXTRACTORS, ".pdf", lambda content: "  extracted from pdf  ")

    assert extract_text("scenario.PDF", b"%PDF-fake") == "extracted from pdf"


def test_content_hash_is_deterministic_sha256():
    digest = content_hash(b"same bytes")

    assert digest == content_hash(b"same bytes")
    assert digest != content_hash(b"different bytes")
    assert len(digest) == 64


class FakeMinioClient:
    def __init__(self, bucket_exists: bool = True):
        self._bucket_exists = bucket_exists
        self.made_bucket = None
        self.put_calls = []
        self.removed = []

    def bucket_exists(self, bucket):
        return self._bucket_exists

    def make_bucket(self, bucket):
        self.made_bucket = bucket

    def put_object(self, bucket, key, stream, length):
        self.put_calls.append((bucket, key, stream.read(), length))

    def remove_object(self, bucket, key):
        self.removed.append((bucket, key))


def test_store_upload_keys_by_content_hash_and_creates_bucket_if_missing(monkeypatch):
    fake = FakeMinioClient(bucket_exists=False)
    monkeypatch.setattr(intake, "_client", lambda: fake)

    stored = intake.store_upload("scenario.pdf", b"pdf bytes")

    assert stored.hash == content_hash(b"pdf bytes")
    assert stored.key == f"uploads/{stored.hash}.pdf"
    assert fake.made_bucket == intake.settings.minio_bucket
    assert fake.put_calls == [(intake.settings.minio_bucket, stored.key, b"pdf bytes", 9)]


def test_store_upload_skips_make_bucket_when_it_exists(monkeypatch):
    fake = FakeMinioClient(bucket_exists=True)
    monkeypatch.setattr(intake, "_client", lambda: fake)

    intake.store_upload("scenario.txt", b"text")

    assert fake.made_bucket is None


def test_delete_upload_removes_the_object(monkeypatch):
    fake = FakeMinioClient()
    monkeypatch.setattr(intake, "_client", lambda: fake)

    intake.delete_upload("uploads/abc.pdf")

    assert fake.removed == [(intake.settings.minio_bucket, "uploads/abc.pdf")]
